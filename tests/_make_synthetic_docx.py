#!/usr/bin/env python3
"""Generate the synthetic docx fixtures committed under tests/fixtures/inputs/.

Two scenarios:

- synthetic_happy.docx — three real, Crossref-indexed Vancouver-numbered
  citations. The pipeline should classify all three as `valid` and emit
  zero entries in any Stage 4b bucket. README's Quickstart points at this.

- synthetic_mixed.docx — six citations covering all four Stage-4b
  bucket states (valid silent, corrections, replacements, unresolvable):
    1, 2: Real, valid Vancouver citations (controls).
    3, 4: Real papers attributed to the wrong author — title, venue,
       year, and pages all match a real paper but the byline credits
       someone else. Stage 3 → partially_valid; Stage 4b →
       corrections[].
    5: Fully fabricated citation with a plausible-looking DOI that
       does not resolve. Stage 3 → hallucinated; Stage 4 finds a real
       replacement (a published paper on the same topic) → Stage 4b
       replacements[].
    6: Fully fabricated citation in a topic where Stage 4 cannot
       confirm any candidate via Channel A. Stage 3 → hallucinated;
       Stage 4 returns no_suitable_replacement → Stage 4b
       unresolvable[].

Determinism: python-docx writes to a zip archive whose entries carry the
current local time as their mtime. We post-process the zip to zero every
entry's timestamp and sort entries lexicographically, so two runs of this
script produce byte-identical output. Acceptance criterion (a) of the
Phase 3 synthetic-fixtures deliverable.

Usage:
    .venv/bin/python3 tests/_make_synthetic_docx.py
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

try:
    from docx import Document
except ImportError:
    sys.stderr.write("ERROR: python-docx not installed. Run: pip install python-docx\n")
    sys.exit(2)

OUTPUT_DIR = Path(__file__).resolve().parent / "fixtures" / "inputs"

# Earliest timestamp the zip format can encode. Using this guarantees
# determinism across runs and across hosts in different timezones.
DETERMINISTIC_DATE = (1980, 1, 1, 0, 0, 0)

HAPPY_TITLE = "Synthetic happy-path document for citation-verify tests"
HAPPY_INTRO = (
    "This document is a deterministic test fixture committed to the citation-verify "
    "repository. The bibliography below contains three citations to widely-cited, "
    "Crossref-indexed papers. The pipeline should classify all three as `valid` and "
    "produce empty corrections, replacements, and unresolvable buckets in Stage 4b."
)
HAPPY_REFS = [
    "1. Vaswani A, Shazeer N, Parmar N, Uszkoreit J, Jones L, Gomez AN, et al. "
    "Attention is all you need. Adv Neural Inf Process Syst. 2017;30:5998-6008.",

    "2. He K, Zhang X, Ren S, Sun J. Deep residual learning for image recognition. "
    "In: Proceedings of the IEEE Conference on Computer Vision and Pattern "
    "Recognition (CVPR). 2016. p. 770-778.",

    "3. LeCun Y, Bengio Y, Hinton G. Deep learning. Nature. 2015;521(7553):436-444.",
]

MIXED_TITLE = "Synthetic mixed-shape document for citation-verify tests"
MIXED_INTRO = (
    "This document exercises every Stage 4b output bucket. Citations 1 and 2 are "
    "clean valid references. Citations 3 and 4 each show an author-attribution "
    "drift — title, venue, year, and pages all match a real paper, but the byline "
    "credits the wrong author — and should land in `corrections[]` after Stage 4b. "
    "Citation 5 is an entirely fabricated paper with a non-resolving DOI in a "
    "topic where a real published paper covers the same claim; Stage 3 labels it "
    "`hallucinated`, Stage 4 finds the real paper, and Stage 4b emits a "
    "`replacements[]` entry. Citation 6 is also fabricated but in a meta-domain "
    "(\"citation hallucination detection\") where Stage 4 cannot confirm any "
    "candidate via Channel A; Stage 4b lands it in `unresolvable[]`."
)
MIXED_REFS = [
    # Citation 1: valid Vancouver — Jinek 2012 (CRISPR-Cas9 discovery).
    "1. Jinek M, Chylinski K, Fonfara I, Hauer M, Doudna JA, Charpentier E. "
    "A programmable dual-RNA-guided DNA endonuclease in adaptive bacterial "
    "immunity. Science. 2012;337(6096):816-821.",

    # Citation 2: valid Vancouver — Takahashi & Yamanaka 2006 (induced pluripotency).
    "2. Takahashi K, Yamanaka S. Induction of pluripotent stem cells from mouse "
    "embryonic and adult fibroblast cultures by defined factors. Cell. "
    "2006;126(4):663-676.",

    # Citation 3: wrong-author drift over Mali et al. 2013 RNA-guided genome editing.
    "3. Smith P, Jones K. RNA-guided human genome engineering via Cas9. Science. "
    "2013;339(6121):823-826.",

    # Citation 4: wrong-author drift over Wang et al. 2014 CRISPR genetic screens.
    "4. Patel R, Kumar S. Genetic screens in human cells using the CRISPR-Cas9 "
    "system. Science. 2014;343(6166):80-84.",

    # Citation 5: fully fabricated; topic exists (off-target prediction) so Stage 4
    # can find a real replacement (Doench et al. 2016 nbt.3437).
    "5. Lee J, Wang K. Universal off-target prediction model for CRISPR-Cas9 "
    "systems. Nature Biotechnology. 2020;38:1234-1242. "
    "doi:10.1038/nbt.fakefake2020",

    # Citation 6: fully fabricated meta-domain claim with no real paper to anchor.
    "6. Anonymous. A novel framework for citation hallucination detection in "
    "academic literature. Nature Methods. 2024;21:9999-10001. "
    "doi:10.1038/nmeth.fakefake2024",
]


def _build_docx_bytes(title: str, intro: str, refs: list[str]) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(intro)
    doc.add_heading("References", level=1)
    for ref in refs:
        doc.add_paragraph(ref)
    raw = io.BytesIO()
    doc.save(raw)
    return _normalize_zip_envelope(raw.getvalue())


def _normalize_zip_envelope(blob: bytes) -> bytes:
    """Re-pack the docx zip with sorted entries and zeroed timestamps so the
    output is byte-identical across runs."""
    src = io.BytesIO(blob)
    dst = io.BytesIO()
    with zipfile.ZipFile(src, "r") as zin:
        infos = sorted(zin.infolist(), key=lambda i: i.filename)
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in infos:
                payload = zin.read(info.filename)
                fresh = zipfile.ZipInfo(info.filename, date_time=DETERMINISTIC_DATE)
                fresh.compress_type = zipfile.ZIP_DEFLATED
                # Strip MS-DOS dirty-bit + any platform-specific external attrs;
                # keep only the file/dir mode bits, normalized to 0o644.
                fresh.external_attr = (0o644 & 0xFFFF) << 16
                fresh.create_system = 0  # MS-DOS / FAT — same on every host
                zout.writestr(fresh, payload)
    return dst.getvalue()


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    happy = _build_docx_bytes(HAPPY_TITLE, HAPPY_INTRO, HAPPY_REFS)
    mixed = _build_docx_bytes(MIXED_TITLE, MIXED_INTRO, MIXED_REFS)

    happy_path = OUTPUT_DIR / "synthetic_happy.docx"
    mixed_path = OUTPUT_DIR / "synthetic_mixed.docx"
    happy_path.write_bytes(happy)
    mixed_path.write_bytes(mixed)

    print(f"wrote {happy_path}  ({len(happy):,} bytes)")
    print(f"wrote {mixed_path}  ({len(mixed):,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
