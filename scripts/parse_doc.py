#!/usr/bin/env python3
"""
parse_doc.py — Extract the references section text (and in-text context) from a PDF or DOCX.

Output JSON to stdout (or --out FILE) with shape:
  {
    "source": "<path>",
    "doc_type": "pdf" | "docx",
    "full_text": "<entire document plain text>",
    "references_text": "<just the bibliography section, raw>",
    "warnings": ["..."]
  }

The heavy lifting (turning the raw bibliography text into structured citations) is done
by one atomic codex-exec call downstream, not here. Keep this module simple and robust.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

# Regex patterns for finding the "References" header. Case-insensitive.
# We accept common variants: "References", "Bibliography", "Works Cited", "参考文献" (rare — doc is English-primary).
REFERENCE_HEADER_PATTERNS = [
    r"^\s*references\s*$",
    r"^\s*bibliography\s*$",
    r"^\s*works\s+cited\s*$",
    r"^\s*literature\s+cited\s*$",
    r"^\s*(?:主要)?参考文献\s*$",            # "参考文献" or "主要参考文献"
    r"^\s*引用文献\s*$",
    r"^\s*参考资料\s*$",
    r"^\s*\d+\.\s+references\s*$",          # "7. References"
    r"^\s*references\s+and\s+notes\s*$",
]

# Text that signals the end of a references section (acknowledgments, appendix, etc.)
END_OF_REFERENCES_PATTERNS = [
    r"^\s*appendix\s*[a-z0-9]*\s*$",
    r"^\s*acknowledgements?\s*$",
    r"^\s*acknowledgments?\s*$",
    r"^\s*supplementary\s+(material|information)\s*$",
    r"^\s*supporting\s+information\s*$",
    r"^\s*author\s+contributions\s*$",
    r"^\s*funding\s*$",
    r"^\s*conflict\s+of\s+interest\s*$",
    r"^\s*declarations?\s*$",
]


def _extract_text_pdf(path: Path) -> str:
    """Extract plain text from a PDF using pdfplumber, one page at a time, joined by \\n\\n."""
    try:
        import pdfplumber
    except ImportError:
        sys.stderr.write(
            "ERROR: pdfplumber not installed. Run: pip install pdfplumber\n"
        )
        sys.exit(2)

    pages = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            pages.append(t)
    return "\n\n".join(pages)


def _extract_text_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        sys.stderr.write(
            "ERROR: python-docx not installed. Run: pip install python-docx\n"
        )
        sys.exit(2)

    document = docx.Document(str(path))
    parts = []
    for p in document.paragraphs:
        if p.text:
            parts.append(p.text)
    # Also include tables (some authors put refs in tables — rare but happens)
    for tbl in document.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _find_references_section(full_text: str) -> tuple[str, list[str]]:
    """Locate the references section. Returns (references_text, warnings)."""
    warnings: list[str] = []
    lines = full_text.splitlines()
    # Find start: last occurrence of a reference header (some papers have "See references" early on)
    start_idx = None
    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue
        # Only check lines that are short-ish — headers are usually on their own line
        if len(stripped) > 60:
            continue
        low = stripped.lower()
        for pat in REFERENCE_HEADER_PATTERNS:
            if re.match(pat, low):
                start_idx = i
                break
    if start_idx is None:
        # Heuristic fallback: take last 25% of the document
        warnings.append("No explicit References header found; falling back to last 25% of document.")
        cutoff = int(len(lines) * 0.75)
        return "\n".join(lines[cutoff:]), warnings

    # Find end: the next end-of-references marker after start_idx
    end_idx = len(lines)
    for j in range(start_idx + 1, len(lines)):
        stripped = lines[j].strip()
        if len(stripped) > 60:
            continue
        low = stripped.lower()
        for pat in END_OF_REFERENCES_PATTERNS:
            if re.match(pat, low):
                end_idx = j
                break
        if end_idx != len(lines):
            break

    ref_text = "\n".join(lines[start_idx + 1 : end_idx]).strip()
    if not ref_text:
        warnings.append(f"References header at line {start_idx} but section body is empty.")
    return ref_text, warnings


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input", help="Path to .pdf or .docx file")
    ap.add_argument("--out", help="Output JSON file (default: stdout)")
    args = ap.parse_args()

    path = Path(args.input)
    if not path.exists():
        sys.stderr.write(f"ERROR: file not found: {path}\n")
        return 1

    ext = path.suffix.lower()
    if ext == ".pdf":
        full_text = _extract_text_pdf(path)
        doc_type = "pdf"
    elif ext in (".docx", ".doc"):
        full_text = _extract_text_docx(path)
        doc_type = "docx"
    else:
        sys.stderr.write(f"ERROR: unsupported file type '{ext}'. Use .pdf or .docx.\n")
        return 1

    ref_text, warnings = _find_references_section(full_text)

    payload = {
        "source": str(path),
        "doc_type": doc_type,
        "full_text": full_text,
        "references_text": ref_text,
        "warnings": warnings,
    }
    out_str = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.out:
        Path(args.out).write_text(out_str, encoding="utf-8")
    else:
        sys.stdout.write(out_str)
    return 0


if __name__ == "__main__":
    sys.exit(main())
